import csv
import io
import json
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter

from munipal.services.register_engine import build_package, csv_bytes, digest, extract_document


def test_exact_excerpts_and_no_invented_dates():
    source = "Section 4.\nThe Borrower shall deliver annual reports within 180 days.\n\nSection 5.\nThe Trustee shall provide notice."
    content = source.encode()
    package, count = build_package(
        "A real input",
        [
            {
                "id": "doc",
                "filename": "CDA.txt",
                "sha256": digest(content),
                "extracted": extract_document("CDA.txt", content),
            }
        ],
        1,
    )
    with zipfile.ZipFile(io.BytesIO(package)) as archive:
        rows = json.loads(archive.read("candidate-map.json"))
        assert count == 2
        assert all(row["quote"] in source for row in rows)
        assert all(row["frequency"] == row["due_rule_text_verbatim"] == "" for row in rows)
        assert all("status" not in row and "due_date" not in row for row in rows)
        assert "may bind another party" in json.loads(archive.read("manifest.json"))["limitation"]
        assert "SYN-HSG" not in archive.read("candidate-map.csv").decode()


def test_formula_safe_csv():
    result = list(
        csv.DictReader(io.StringIO(csv_bytes([{"quote": " =HYPERLINK(1)"}]).decode("utf-8-sig")))
    )
    assert result[0]["quote"].startswith("'")


def test_docx_paragraphs_and_tables():
    document = Document()
    document.add_paragraph("The Borrower shall deliver reports.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "The Borrower shall maintain insurance."
    buffer = io.BytesIO()
    document.save(buffer)
    parts = extract_document("deal.docx", buffer.getvalue())
    assert len(parts) == 2 and parts[1]["locator"] == "Table 1, row 1"


def test_blank_pdf_requests_ocr():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(ValueError, match="OCR"):
        extract_document("scan.pdf", buffer.getvalue())


@pytest.mark.parametrize(
    "filename,content", [("bad.exe", b"123"), ("blank.txt", b""), ("binary.txt", b"\x00text")]
)
def test_rejects_unsupported_or_unreadable_documents(filename, content):
    with pytest.raises(ValueError):
        extract_document(filename, content)
