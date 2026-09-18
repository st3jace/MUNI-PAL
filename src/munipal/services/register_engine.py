"""Production candidate packaging derived from fulfillment/demo/run.py duty matching.

No demo paths, fixture approvals, LLM calls, date arithmetic or status assignments.
Results are a bounded lexical first pass, not the professionally approved register.
"""

import csv
import hashlib
import io
import json
import re
import zipfile
from importlib.resources import files
from pathlib import PurePosixPath

from docx import Document
from pypdf import PdfReader

DUTY_RE = re.compile(
    r"\b(?:shall|must|agrees to)\s+(?:deliver|provide|give|furnish|notify|retain|obtain|cause|send|maintain|pay|file|certify|compute)\b",
    re.I,
)
SECTION_RE = re.compile(r"(?:Section|ARTICLE)\s+[0-9A-Za-z.()]+", re.I)
FIELDS = [
    "id",
    "obligation",
    "source_document",
    "clause",
    "recipient",
    "frequency",
    "due_rule_text_verbatim",
    "quote",
    "evidence_expected",
    "counsel_question",
    "source_locator",
    "source_sha256",
]
LIMITATION = (
    "This is a candidate map, not an approved obligation register. It is a lexical first pass "
    "over extracted text and may miss duties, split clauses, tables or scanned pages. "
    "A matching clause may bind another party. Verify each source and have your bond counsel "
    "or dissemination agent approve the list, responsible party and dates in writing. "
    "Blank fields are awaiting input; no dates, filing statuses or compliance conclusions "
    "are generated. No reminders, automation jobs or filings have been activated."
)


def extract_document(filename: str, content: bytes) -> list[dict]:
    """Extract bounded text with locators. Reject unreadable documents explicitly."""
    ext = PurePosixPath(filename).suffix.lower()
    parts = []
    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted or len(reader.pages) > 300:
            raise ValueError("Use an unencrypted PDF of at most 300 pages.")
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if not text.strip():
                raise ValueError(f"Page {i} has no extractable text. Upload an OCR text version.")
            parts.append({"locator": f"Page {i}", "text": text})
    elif ext == ".docx":
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 20_000_000:
                raise ValueError("The expanded document exceeds 20 MB.")
        doc = Document(io.BytesIO(content))
        parts = [
            {"locator": f"Paragraph {i}", "text": p.text}
            for i, p in enumerate(doc.paragraphs, 1)
            if p.text.strip()
        ]
        for table_num, table in enumerate(doc.tables, 1):
            for row_num, row in enumerate(table.rows, 1):
                parts.append(
                    {
                        "locator": f"Table {table_num}, row {row_num}",
                        "text": " | ".join(c.text for c in row.cells),
                    }
                )
    elif ext in {".txt", ".md"}:
        text = content.decode("utf-8-sig")
        if "\x00" in text:
            raise ValueError("Use a UTF-8 text document.")
        parts = [{"locator": "Full text (see paragraph locator)", "text": text}]
    else:
        raise ValueError("Upload PDF, DOCX, TXT or Markdown documents.")
    if not parts or not any(p["text"].strip() for p in parts):
        raise ValueError("No text found. Upload a searchable document or OCR text.")
    if sum(len(p["text"]) for p in parts) > 1_000_000:
        raise ValueError("Extracted text exceeds one million characters. Split this document.")
    return parts


def csv_bytes(rows: list[dict]) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=FIELDS)
    writer.writeheader()
    # Spreadsheet formula protection applies to every untrusted cell.
    for row in rows:
        writer.writerow(
            {
                k: "'" + str(v) if str(v).lstrip().startswith(("=", "+", "-", "@")) else v
                for k, v in row.items()
            }
        )
    return output.getvalue().encode("utf-8-sig")


def build_package(deal_name: str, documents: list[dict], version: int) -> tuple[bytes, int]:
    candidates = []
    for document in documents:
        for part in document["extracted"]:
            # Keep contiguous excerpts intact. Locate headings but do not infer parties or dates.
            paragraphs = re.split(r"\n\s*\n", part["text"])
            for i, paragraph in enumerate(paragraphs, 1):
                if not DUTY_RE.search(paragraph):
                    continue
                section = SECTION_RE.search(paragraph)
                row = dict.fromkeys(FIELDS, "")
                row.update(
                    id=f"OB-{len(candidates) + 1:03d}",
                    source_document=document["filename"],
                    clause=section.group(0) if section else "",
                    quote=paragraph,
                    source_locator=f"{part['locator']}, paragraph {i}",
                    source_sha256=document["sha256"],
                    counsel_question="Does this clause bind the obligated person? Confirm the deliverables, recipient, frequency and approved dates in writing.",
                )
                candidates.append(row)
    manifest = {
        "deal": deal_name,
        "version": version,
        "engine": "candidate-lexical-v1",
        "limitation": LIMITATION,
        "candidate_count": len(candidates),
        "documents": [{k: d[k] for k in ("id", "filename", "sha256")} for d in documents],
    }
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("candidate-map.csv", csv_bytes(candidates))
        archive.writestr("candidate-map.json", json.dumps(candidates, indent=2))
        archive.writestr("manifest.json", json.dumps(manifest, indent=2))
        archive.writestr(
            "READ-ME.md",
            f"# {deal_name}\n\n{LIMITATION}\n\n"
            "Use the enclosed skill with your own AI and original documents for a fuller draft. "
            "Take the candidate map to your professional team. Import only their approved dates "
            "into your own calendar and automation tools.\n",
        )
        for name in ("SKILL.md", "CHECKLIST.md", "LEGEND.md", "REGISTER-TEMPLATE.md"):
            archive.writestr(
                name, files("munipal.services").joinpath("register_kit", name).read_bytes()
            )
    return buffer.getvalue(), len(candidates)


def digest(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()
